"""Preview Service — Converts PPTX, DOCX, and PDF documents into cached PDF/Slide PNG preview assets."""

from __future__ import annotations

import json
import os
import fitz  # PyMuPDF
from pathlib import Path
from typing import Any, Dict, List, Optional
from datetime import datetime, timezone

from app.core.config import settings
from app.core.logging import get_logger
from app.models.database import get_document_by_id
from app.utils.hashing import hash_file

log = get_logger(__name__)

PREVIEWS_DIR = Path(settings.DB_PATH).parent / "previews"
PREVIEWS_DIR.mkdir(parents=True, exist_ok=True)


def get_or_create_preview(doc_id: str, force: bool = False) -> Dict[str, Any]:
    """Retrieve existing preview metadata or generate a new conversion cache."""
    doc = get_document_by_id(doc_id)
    if not doc:
        raise ValueError(f"Document not found: {doc_id}")

    file_path = Path(doc["file_path"])
    if not file_path.exists():
        raise FileNotFoundError(f"File missing on server: {file_path}")

    current_hash = hash_file(file_path)
    doc_preview_dir = PREVIEWS_DIR / doc_id
    doc_preview_dir.mkdir(parents=True, exist_ok=True)
    meta_path = doc_preview_dir / "preview.json"

    # Check existing valid cache
    if meta_path.exists() and not force:
        try:
            with open(meta_path, "r", encoding="utf-8") as f:
                meta = json.load(f)
            if meta.get("file_hash") == current_hash and meta.get("status") == "ready":
                return meta
        except Exception as e:
            log.warning("Failed reading preview meta for %s, re-generating: %s", doc_id, e)

    file_type = doc["file_type"].lower()
    log.info("Generating preview assets for document %s (type: %s)", doc_id, file_type)

    meta = {
        "document_id": doc_id,
        "filename": doc["filename"],
        "file_type": file_type,
        "file_hash": current_hash,
        "status": "processing",
        "page_count": 0,
        "preview_type": "unknown",
        "has_pdf": False,
        "has_images": False,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }

    try:
        if file_type == "pdf":
            _generate_pdf_preview(file_path, doc_preview_dir, meta)
        elif file_type in ("pptx", "ppt"):
            _generate_pptx_preview(file_path, doc_preview_dir, meta)
        elif file_type in ("docx", "doc"):
            _generate_docx_preview(file_path, doc_preview_dir, meta)
        elif file_type in ("png", "jpg", "jpeg", "webp"):
            _generate_image_preview(file_path, doc_preview_dir, meta)
        else:
            _generate_text_preview(file_path, doc_preview_dir, meta)

        meta["status"] = "ready"
    except Exception as e:
        log.error("Failed preview generation for %s: %s", doc_id, e)
        meta["status"] = "failed"
        meta["error"] = str(e)

    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2)

    return meta


def get_preview_pdf_path(doc_id: str) -> Optional[Path]:
    """Get path to preview PDF if available."""
    pdf_file = PREVIEWS_DIR / doc_id / "preview.pdf"
    return pdf_file if pdf_file.exists() else None


def get_slide_image_path(doc_id: str, slide_num: int) -> Optional[Path]:
    """Get path to a specific slide PNG image (1-indexed)."""
    img_file = PREVIEWS_DIR / doc_id / f"slide-{slide_num:03d}.png"
    return img_file if img_file.exists() else None


# ── Conversion Strategy Implementations ──────────────────────────

def _generate_pdf_preview(pdf_path: Path, out_dir: Path, meta: Dict[str, Any]) -> None:
    """Extract page count and render page PNGs for native PDF."""
    target_pdf = out_dir / "preview.pdf"
    if pdf_path.resolve() != target_pdf.resolve():
        import shutil
        shutil.copy2(pdf_path, target_pdf)

    doc = fitz.open(str(target_pdf))
    meta["page_count"] = len(doc)
    meta["preview_type"] = "pdf_and_images"
    meta["has_pdf"] = True

    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=150)
        img_path = out_dir / f"slide-{(i + 1):03d}.png"
        pix.save(str(img_path))

    doc.close()
    meta["has_images"] = True


def _generate_pptx_preview(pptx_path: Path, out_dir: Path, meta: Dict[str, Any]) -> None:
    """Convert PPTX to PDF using PowerPoint COM (Windows) or python-pptx fallback."""
    target_pdf = out_dir / "preview.pdf"
    converted_pdf = False

    # Primary Strategy: PowerPoint COM on Windows
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()

        ppt_app = win32com.client.DispatchEx("PowerPoint.Application")
        # ppWindowStateMinimized = 2
        try:
            presentation = ppt_app.Presentations.Open(
                str(pptx_path.resolve()), WithWindow=False, ReadOnly=True
            )
            # Format 32 = ppSaveAsPDF
            presentation.SaveAs(str(target_pdf.resolve()), 32)
            presentation.Close()
            converted_pdf = True
            log.info("✓ Converted PPTX → PDF via PowerPoint COM: %s", pptx_path.name)
        finally:
            ppt_app.Quit()
            pythoncom.CoUninitialize()
    except Exception as com_err:
        log.warning("PowerPoint COM conversion failed for %s: %s", pptx_path.name, com_err)

    if converted_pdf and target_pdf.exists():
        _generate_pdf_preview(target_pdf, out_dir, meta)
        meta["preview_type"] = "pptx_pdf"
        return

    # Secondary Strategy: python-pptx fallback to render slides visually as layout cards
    _fallback_pptx_render(pptx_path, out_dir, meta)


def _generate_docx_preview(docx_path: Path, out_dir: Path, meta: Dict[str, Any]) -> None:
    """Convert DOCX to PDF using Word COM (Windows) or fallback."""
    target_pdf = out_dir / "preview.pdf"
    converted_pdf = False

    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()

        word_app = win32com.client.DispatchEx("Word.Application")
        word_app.Visible = False
        try:
            doc = word_app.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
            # Format 17 = wdFormatPDF
            doc.SaveAs(str(target_pdf.resolve()), 17)
            doc.Close(False)
            converted_pdf = True
            log.info("✓ Converted DOCX → PDF via Word COM: %s", docx_path.name)
        finally:
            word_app.Quit()
            pythoncom.CoUninitialize()
    except Exception as com_err:
        log.warning("Word COM conversion failed for %s: %s", docx_path.name, com_err)

    if converted_pdf and target_pdf.exists():
        _generate_pdf_preview(target_pdf, out_dir, meta)
        meta["preview_type"] = "docx_pdf"
        return

    _fallback_text_render(docx_path, out_dir, meta)


def _generate_image_preview(img_path: Path, out_dir: Path, meta: Dict[str, Any]) -> None:
    """Single image preview."""
    import shutil
    target_img = out_dir / "slide-001.png"
    shutil.copy2(img_path, target_img)
    meta["page_count"] = 1
    meta["preview_type"] = "image"
    meta["has_pdf"] = False
    meta["has_images"] = True


def _generate_text_preview(txt_path: Path, out_dir: Path, meta: Dict[str, Any]) -> None:
    """Text document preview."""
    meta["page_count"] = 1
    meta["preview_type"] = "text"
    meta["has_pdf"] = False
    meta["has_images"] = False


def _fallback_pptx_render(pptx_path: Path, out_dir: Path, meta: Dict[str, Any]) -> None:
    """Render slide layouts as PNG images using PIL if COM is not available."""
    try:
        from pptx import Presentation
        from PIL import Image, ImageDraw, ImageFont

        prs = Presentation(str(pptx_path))
        num_slides = len(prs.slides)
        meta["page_count"] = num_slides
        meta["preview_type"] = "pptx_fallback"
        meta["has_pdf"] = False

        for idx, slide in enumerate(prs.slides):
            # 16:9 Slide Canvas (1280x720)
            img = Image.new("RGB", (1280, 720), color=(248, 249, 250))
            draw = ImageDraw.Draw(img)

            # Slide Header
            draw.rectangle([(0, 0), (1280, 70)], fill=(30, 30, 30))
            draw.text((30, 20), f"SLIDE {idx + 1} / {num_slides} — {pptx_path.name}", fill=(255, 255, 255))

            # Extract text elements
            y_offset = 100
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        txt = paragraph.text.strip()
                        if txt:
                            draw.text((50, y_offset), txt[:120], fill=(20, 20, 20))
                            y_offset += 35
                            if y_offset > 650:
                                break

            img_path = out_dir / f"slide-{(idx + 1):03d}.png"
            img.save(str(img_path))

        meta["has_images"] = True
        log.info("✓ Rendered %d PPTX slides via Pillow fallback", num_slides)
    except Exception as fallback_err:
        log.error("PPTX fallback rendering failed: %s", fallback_err)
        meta["page_count"] = 1
        meta["preview_type"] = "failed"


def _fallback_text_render(doc_path: Path, out_dir: Path, meta: Dict[str, Any]) -> None:
    meta["page_count"] = 1
    meta["preview_type"] = "text"
