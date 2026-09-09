"""DOCX parser — extracts text from Word documents."""

from __future__ import annotations

import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Tuple

from app.core.logging import get_logger
from app.utils.text import clean_text

log = get_logger(__name__)


def _raw_xml_parse(filepath: Path) -> List[Tuple[str, Dict[str, Any]]]:
    """Fallback XML parser for docx files."""
    try:
        with zipfile.ZipFile(filepath) as z:
            xml_content = z.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            texts = [node.text for node in tree.iter() if node.tag.endswith("}t") and node.text]
            full_text = "\n".join(texts)
            cleaned = clean_text(full_text)
            if cleaned:
                return [(cleaned, {"section": "Document Content", "source_type": "docx"})]
    except Exception as e:
        log.error("DOCX raw XML fallback error for %s: %s", filepath.name, e)
    return []


def parse(filepath: str | Path) -> List[Tuple[str, Dict[str, Any]]]:
    filepath = Path(filepath)
    results: List[Tuple[str, Dict[str, Any]]] = []

    try:
        from docx import Document
        doc = Document(str(filepath))
        sections: List[Tuple[str, str]] = []
        current_section = ""
        current_heading = ""

        for para in doc.paragraphs:
            if para.style and para.style.name and para.style.name.startswith("Heading"):
                if current_section.strip():
                    sections.append((current_heading, current_section))
                current_heading = para.text.strip()
                current_section = para.text + "\n"
            else:
                current_section += para.text + "\n"

        if current_section.strip():
            sections.append((current_heading, current_section))

        # Extract tables
        for i, table in enumerate(doc.tables):
            table_text = ""
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    table_text += " | ".join(cells) + "\n"
            if table_text.strip():
                sections.append((f"Table {i+1}", "[TABLE]\n" + table_text))

        for heading, text in sections:
            cleaned = clean_text(text)
            if cleaned:
                results.append((cleaned, {
                    "section": heading,
                    "source_type": "docx",
                }))

    except Exception as e:
        log.warning("DOCX python-docx parse warning for %s: %s — attempting XML fallback", filepath.name, e)
        results = _raw_xml_parse(filepath)

    if not results:
        results = _raw_xml_parse(filepath)

    log.info("DOCX parsed: %s → %d sections", filepath.name, len(results))
    return results
