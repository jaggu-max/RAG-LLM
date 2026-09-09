"""PPTX parser — extracts text slide-by-slide."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Tuple

from app.core.logging import get_logger
from app.utils.text import clean_text

log = get_logger(__name__)


def parse(filepath: str | Path) -> List[Tuple[str, Dict[str, Any]]]:
    from pptx import Presentation

    filepath = Path(filepath)
    results: List[Tuple[str, Dict[str, Any]]] = []

    try:
        prs = Presentation(str(filepath))
        for slide_num, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            texts.append(t)
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            texts.append(" | ".join(cells))

            text = "\n".join(texts)
            cleaned = clean_text(text)
            if cleaned:
                results.append((cleaned, {
                    "slide_number": slide_num,
                    "source_type": "pptx",
                }))
    except Exception as e:
        log.error("PPTX parse error for %s: %s", filepath.name, e)
        raise

    log.info("PPTX parsed: %s → %d slides", filepath.name, len(results))
    return results
